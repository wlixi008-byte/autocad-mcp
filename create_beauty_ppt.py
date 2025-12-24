import collections 
import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def create_presentation():
    prs = Presentation()

    # Define a simple function to add a slide with title and content
    def add_slide(title_text, content_text_list, layout_index=1):
        """
        layout_index: 0=Title, 1=Title+Content, ...
        """
        slide_layout = prs.slide_layouts[layout_index]
        slide = prs.slides.add_slide(slide_layout)
        
        # Set Title
        if slide.shapes.title:
            slide.shapes.title.text = title_text
            
        # Set Content (if available)
        if layout_index == 1: # Title + Content layout
            # Find the content placeholder
            for shape in slide.placeholders:
                if shape.placeholder_format.idx == 1: # Usually body
                    text_frame = shape.text_frame
                    text_frame.clear()  # Clear existing
                    
                    for i, line in enumerate(content_text_list):
                        p = text_frame.add_paragraph()
                        p.text = line
                        p.font.size = Pt(20)
                        if i == 0:
                            p.font.bold = True
                        else:
                            p.level = 0 # Bullet level

    # --- Slide 1: Title Slide ---
    slide_layout = prs.slide_layouts[0] # Title Slide
    slide = prs.slides.add_slide(slide_layout)
    if slide.shapes.title:
        slide.shapes.title.text = "建筑设计审美提升"
    if slide.placeholders[1]: # Subtitle
        slide.placeholders[1].text = "从重复劳动到算法美学 - 基于 MCP Skills 的新范式"

    # --- Slide 2: Core Concept ---
    add_slide(
        "什么是 Skills?",
        [
            "定义：",
            "Skills 是将建筑设计规则转化为可复用的代码模块。",
            "",
            "核心价值：",
            "- 将重复性绘图工作自动化 (Python 脚本)。",
            "- 通过 MCP (Model Context Protocol) 连接 AI 与 CAD 工具。",
            "- 让设计师像管理资产一样管理'设计能力'。"
        ]
    )

    # --- Slide 3: The Pain Point ---
    add_slide(
        "传统设计的审美困境",
        [
            "现状：",
            "1. 时间被机械绘图占据 (画墙、标尺寸、改图层)。",
            "2. 人为误差导致图面凌乱 (线头未闭合、标注重叠)。",
            "3. 修改成本高，挤压了推敲比例与美学的时间。",
            "",
            "后果：",
            "设计师沦为'绘图员'，审美创造力无法充分释放。"
        ]
    )

    # --- Slide 4: Automation & Aesthetics ---
    add_slide(
        "Skills 如何提升审美?",
        [
            "1. 绝对的精确性",
            "- 算法生成的几何图形（如户型图）绝对规整。",
            "- 墙体厚度、倒角、对齐由代码控制，杜绝'手抖'。",
            "",
            "2. 统一的视觉标准",
            "- 自动应用标准字体、字号和图层颜色。",
            "- 确保整套图纸风格高度一致，通过整洁体现专业美感。"
        ]
    )

    # --- Slide 5: Efficiency -> Creativity ---
    add_slide(
        "效率释放创造力",
        [
            "当绘图不再耗时：",
            "- 设计师可以用 80% 的时间思考空间关系、材质与光影。",
            "- 快速生成多种方案进行对比 (Generative Design)。",
            "- 从'如何画出来' 转向 '设计什么更好看'。",
            "",
            "MCP 的角色：",
            "作为 AI 助理，实时响应指令，瞬间完成枯燥工作。"
        ]
    )

    # --- Slide 6: Case Study ---
    add_slide(
        "案例：自动户型生成",
        [
            "输入：简单指令 '画一个 10x11m 的户型'。",
            "输出：",
            "- 完美的墙体结构 (draw_specific_plan.py)。",
            "- 自动布局楼梯、厨房、卫生间。",
            "- 自动添加标准化的中文标注。",
            "",
            "结果：",
            "无需人工干预，直接得到一张符合制图规范、美观清晰的底图。"
        ]
    )

    # --- Slide 7: Conclusion ---
    add_slide(
        "结语：重塑设计流程",
        [
            "未来已来：",
            "Skills + MCP 不仅仅是工具的升级，更是设计思维的解放。",
            "",
            "愿景：",
            "让算法处理繁琐，让人类回归审美。",
            "共建高质量、高效率的'人机协作'建筑设计新时代。"
        ]
    )

    output_path = "Architectural_Aesthetics.pptx"
    prs.save(output_path)
    print(f"Presentation saved to {output_path}")

if __name__ == "__main__":
    create_presentation()
